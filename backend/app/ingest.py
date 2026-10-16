"""Stage 1-3 of the pipeline: extract text, chunk it, and embed it.

Extraction keeps page numbers so answers can cite their source page.
Chunking is paragraph-aware with overlap so context is not cut mid-thought.
"""
import io
import re

import docx
import numpy as np
from pypdf import PdfReader

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}
CHUNK_SIZE = 900
CHUNK_OVERLAP = 150


def extract_pages(filename: str, data: bytes) -> list[tuple[int, str]]:
    """Return a list of (page_number, text) for the uploaded file."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in SUPPORTED:
        raise ValueError(f"Unsupported file type '{ext or 'none'}'. Use PDF, DOCX, TXT or MD.")

    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        pages = [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
        if not any(text.strip() for _, text in pages):
            raise ValueError("No selectable text found. Scanned PDFs need OCR first.")
        return pages

    if ext == ".docx":
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return [(1, "\n".join(parts))]

    return [(1, data.decode("utf-8", errors="replace"))]


def _split_long(paragraph: str, size: int) -> list[str]:
    """Hard-split a paragraph that is longer than one chunk, on word boundaries."""
    pieces, current = [], ""
    for word in paragraph.split():
        if len(current) + len(word) + 1 > size and current:
            pieces.append(current)
            current = word
        else:
            current = f"{current} {word}".strip()
    if current:
        pieces.append(current)
    return pieces


def chunk_pages(pages: list[tuple[int, str]], size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[dict]:
    chunks: list[dict] = []
    for page_no, text in pages:
        paragraphs = []
        for raw in re.split(r"\n\s*\n|\n", text):
            raw = raw.strip()
            if raw:
                paragraphs.extend(_split_long(raw, size) if len(raw) > size else [raw])

        buffer = ""
        for para in paragraphs:
            if buffer and len(buffer) + len(para) + 1 > size:
                chunks.append({"page": page_no, "text": buffer})
                room = size - len(para) - 1  # the carried tail must leave room for the new paragraph
                tail = buffer[-min(overlap, room):] if room > 0 else ""
                tail = tail.split(" ", 1)[-1] if " " in tail else ""  # start the tail on a whole word
                buffer = f"{tail} {para}".strip()
            else:
                buffer = f"{buffer}\n{para}".strip() if buffer else para
        if buffer:
            chunks.append({"page": page_no, "text": buffer})
    return chunks


class HashingEmbedder:
    """Offline embedder: hashed unigrams and bigrams with log-scaled counts.

    Zero API cost and deterministic, so the app runs out of the box. Swap in a
    neural model (sentence-transformers, OpenAI, Voyage) behind the same
    `embed(texts) -> np.ndarray` interface for better semantic recall.
    """

    name = "hashing-v1"
    dim = 512
    _token = re.compile(r"[a-z0-9+#.]+")
    _stop = {
        "the", "a", "an", "and", "or", "of", "to", "in", "for", "on", "with",
        "is", "are", "was", "were", "be", "by", "as", "at", "from", "that", "this", "it",
    }

    def tokens(self, text: str) -> list[str]:
        return [t for t in self._token.findall(text.lower()) if t not in self._stop]

    def embed(self, texts: list[str]) -> np.ndarray:
        matrix = np.zeros((len(texts), self.dim), dtype=np.float32)
        for row, text in enumerate(texts):
            toks = self.tokens(text)
            feats = toks + [f"{a}_{b}" for a, b in zip(toks, toks[1:])]
            counts: dict[int, int] = {}
            for feat in feats:
                idx = _stable_hash(feat) % self.dim
                counts[idx] = counts.get(idx, 0) + 1
            for idx, count in counts.items():
                matrix[row, idx] = 1.0 + np.log(count)
        norms = np.linalg.norm(matrix, axis=1, keepdims=True)
        return matrix / np.maximum(norms, 1e-9)


def _stable_hash(text: str) -> int:
    """FNV-1a so the hash is the same across processes (Python's hash() is salted)."""
    h = 2166136261
    for byte in text.encode("utf-8"):
        h = ((h ^ byte) * 16777619) & 0xFFFFFFFF
    return h
